# src/load_documents.py
from pathlib import Path
import json
from docx import Document
import pymupdf  # 需要 PyMuPDF >= 1.24.3；旧版请改用 import fitz

RAW_DIR = Path("data/raw")
OUT_DIR = Path("data/processed")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def extract_docx(path: Path) -> dict:
    doc = Document(str(path))
    paragraphs = []

    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style = p.style.name if p.style is not None else ""
        paragraphs.append({
            "idx": idx,
            "style": style,
            "text": text,
            "page": None,          # docx 没有页概念，统一为 None
        })

    headers, footers = [], []
    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                headers.append(p.text.strip())
        for p in section.footer.paragraphs:
            if p.text.strip():
                footers.append(p.text.strip())

    return {
        "source_file": path.name,
        "source_type": "docx",
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "headers": headers,
        "footers": footers,
    }


def extract_pdf(path: Path) -> dict:
    """PDF 只有行没有段落，这里按行展开成 paragraphs，保留页码。
    行合并留到 chunk_documents.py 的 merge_pdf_lines() 里做。
    """
    doc = pymupdf.open(str(path))
    paragraphs = []

    for page_idx, page in enumerate(doc, start=1):
        text = page.get_text()
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            paragraphs.append({
                "idx": len(paragraphs),
                "style": "",
                "text": line,
                "page": page_idx,
            })

    doc.close()

    return {
        "source_file": path.name,
        "source_type": "pdf",
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "headers": [],
        "footers": [],
    }


def main():
    # 合并两类文件，避免 return 提前退出
    files = sorted(RAW_DIR.glob("*.docx")) + sorted(RAW_DIR.glob("*.pdf"))
    if not files:
        print(f"[WARN] {RAW_DIR} 下没有 .docx 或 .pdf 文件")
        return

    for path in files:
        suffix = path.suffix.lower()
        if suffix == ".docx":
            data = extract_docx(path)
        elif suffix == ".pdf":
            data = extract_pdf(path)
        else:
            continue

        out_path = OUT_DIR / f"{path.stem}.raw.json"
        out_path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(f"[OK] {path.name} ({data['source_type']}) -> {out_path}  "
              f"({data['paragraph_count']} 个原始块)")


if __name__ == "__main__":
    main()